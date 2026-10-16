"""
Unit and integration tests for latex2pptx.
"""

import os
import subprocess
import tempfile
import unittest
from pptx import Presentation
from PIL import Image

from latex2pptx.cli import build_parser
from latex2pptx.compiler import find_missing_images, create_placeholder_image, compile_latex
from latex2pptx.converter import convert_pdf_to_pptx
from latex2pptx.text import compute_paragraph_font_scale, measure_line_width_pt


class TestLatex2Pptx(unittest.TestCase):
    def test_cli_parser(self):
        parser = build_parser()
        args = parser.parse_args(["clase1.tex", "-o", "output.pptx", "-v"])
        self.assertEqual(args.input, "clase1.tex")
        self.assertEqual(args.output, "output.pptx")
        self.assertTrue(args.verbose)
        self.assertEqual(args.engine, "pdflatex")
        self.assertEqual(args.passes, 2)

    def test_paragraph_font_scale_no_overflow(self):
        sample_line = {
            "x0": 36.5,
            "x1": 305.2,
            "runs": [
                {
                    "family": "Liberation Sans",
                    "bold": False,
                    "italic": False,
                    "size": 10.5,
                    "text": "cada seis meses la autorización escrita de sus afiliados para el",
                }
            ],
        }
        unscaled_w = measure_line_width_pt(sample_line)
        pdf_w = sample_line["x1"] - sample_line["x0"]
        self.assertGreater(unscaled_w, pdf_w)
        scale = compute_paragraph_font_scale([sample_line], "left")
        self.assertLessEqual(unscaled_w * scale, pdf_w + 0.55)

    def test_find_missing_images(self):
        sample_tex = r"""
        \newcommand{\usebg}[1]{\includegraphics{bg#1}}
        \usebg{1}
        \usebg{2}
        \includegraphics[width=2cm]{test_avatar.jpg}
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            missing = find_missing_images(sample_tex, [tmpdir])
            self.assertIn("bg1", missing)
            self.assertIn("bg2", missing)
            self.assertIn("test_avatar.jpg", missing)

            # Create test_avatar.jpg and check that it is found
            open(os.path.join(tmpdir, "test_avatar.jpg"), "w").close()
            missing2 = find_missing_images(sample_tex, [tmpdir])
            self.assertNotIn("test_avatar.jpg", missing2)
            self.assertIn("bg1", missing2)

    def test_placeholder_generator(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            p_bg = create_placeholder_image("bg1", tmpdir)
            self.assertTrue(os.path.isfile(p_bg))
            with Image.open(p_bg) as im:
                self.assertEqual(im.size, (1600, 900))

            p_logo = create_placeholder_image("company_logo.png", tmpdir)
            self.assertTrue(os.path.isfile(p_logo))
            with Image.open(p_logo) as im:
                self.assertEqual(im.size, (400, 150))

    def test_minimal_beamer_compile_and_convert(self):
        min_tex = r"""
\documentclass[aspectratio=169, 11pt]{beamer}
\usepackage{tikz}
\definecolor{myBlue}{HTML}{004B87}
\begin{document}
\begin{frame}
  \frametitle{Test Slide}
  \begin{tikzpicture}
    \fill[myBlue, rounded corners=5pt] (0, 0) rectangle (5, 2);
    \node[text=white, font=\bfseries] at (2.5, 1) {Hello PPTX};
  \end{tikzpicture}
  \only<2->{
    \begin{tikzpicture}
      \fill[red] (6, 0) rectangle (8, 2);
    \end{tikzpicture}
  }
\end{frame}
\end{document}
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            tex_file = os.path.join(tmpdir, "test_min.tex")
            with open(tex_file, "w") as f:
                f.write(min_tex)

            pdf_file = compile_latex(tex_file, allow_placeholders=True, verbose=False)
            self.assertTrue(os.path.isfile(pdf_file))

            pptx_file = os.path.join(tmpdir, "test_min.pptx")
            convert_pdf_to_pptx(pdf_file, pptx_file)
            self.assertTrue(os.path.isfile(pptx_file))

            prs = Presentation(pptx_file)
            # Overlay \only<2-> should result in 2 slides
            self.assertEqual(len(prs.slides), 2)
            # Verify 16:9 dimensions
            self.assertAlmostEqual(prs.slide_width.pt / prs.slide_height.pt, 16.0 / 9.0, places=2)

            # Check that text 'Hello PPTX' was extracted
            found_text = False
            for shape in prs.slides[0].shapes:
                if shape.has_text_frame and "Hello PPTX" in shape.text:
                    found_text = True
                    break
            self.assertTrue(found_text)


if __name__ == "__main__":
    unittest.main()
