"""
Unit and integration tests for markdown speaker notes feature.
"""

import os
import tempfile
import unittest
from pptx import Presentation

from latex2pptx.cli import build_parser
from latex2pptx.compiler import compile_latex
from latex2pptx.converter import convert_pdf_to_pptx
from latex2pptx.notes import (
    parse_markdown_notes,
    load_and_parse_notes,
    apply_notes_to_presentation,
    apply_notes_to_slide,
)


class TestSpeakerNotes(unittest.TestCase):
    def test_cli_parser_notes_arg(self):
        parser = build_parser()
        args = parser.parse_args(["clase1.tex", "-n", "notes.md", "-o", "out.pptx"])
        self.assertEqual(args.notes_file, "notes.md")

    def test_parse_explicit_headings(self):
        md = """
# Slide 1: Welcome
Welcome everyone to the workshop.
- Intro
- Rules

## Slide 2: Agenda
Agenda overview.

### Diapositiva 3 - Objetivos
Objetivos de la sesión.

# 4. Conclusiones
Puntos clave.
"""
        notes = parse_markdown_notes(md)
        self.assertIn(1, notes)
        self.assertIn("## Welcome", notes[1])
        self.assertIn("Welcome everyone", notes[1])

        self.assertIn(2, notes)
        self.assertIn("Agenda overview", notes[2])

        self.assertIn(3, notes)
        self.assertIn("Objetivos de la sesión", notes[3])

        self.assertIn(4, notes)
        self.assertIn("Puntos clave", notes[4])

    def test_parse_slide_ranges(self):
        md = """
# Slide 1-3: Multi-step Overlay
Notes for the three overlay steps.
- Point A
- Point B

## Diapositivas 4 a 6: Casos Prácticos
Casos prácticos explicados.
"""
        notes = parse_markdown_notes(md)
        for s in (1, 2, 3):
            self.assertIn(s, notes)
            self.assertIn("Multi-step Overlay", notes[s])
            self.assertIn("Point A", notes[s])

        for s in (4, 5, 6):
            self.assertIn(s, notes)
            self.assertIn("Casos Prácticos", notes[s])

    def test_parse_comment_directives(self):
        md = """
<!-- slide: 1 -->
Notes for slide 1.

<!-- slide: 2-3: Intro -->
Notes for slides 2 and 3.
"""
        notes = parse_markdown_notes(md)
        self.assertIn(1, notes)
        self.assertEqual(notes[1].strip(), "Notes for slide 1.")
        self.assertIn(2, notes)
        self.assertIn(3, notes)
        self.assertIn("Notes for slides 2 and 3.", notes[2])

    def test_parse_sequential_delimiters(self):
        md = """
First slide notes.
- bullet 1
---
Second slide notes.
---
Third slide notes.
"""
        notes = parse_markdown_notes(md)
        self.assertEqual(len(notes), 3)
        self.assertIn("First slide", notes[1])
        self.assertIn("Second slide", notes[2])
        self.assertIn("Third slide", notes[3])

    def test_apply_notes_to_presentation(self):
        prs = Presentation()
        for _ in range(3):
            prs.slides.add_slide(prs.slide_layouts[6])

        notes_map = {
            1: "Speaker notes for slide 1.",
            3: "Speaker notes for slide 3.",
        }
        applied = apply_notes_to_presentation(prs, notes_map)
        self.assertEqual(applied, 2)

        self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, "Speaker notes for slide 1.")
        self.assertEqual(prs.slides[2].notes_slide.notes_text_frame.text, "Speaker notes for slide 3.")

    def test_end_to_end_beamer_with_notes(self):
        min_tex = r"""
\documentclass[aspectratio=169, 11pt]{beamer}
\begin{document}
\begin{frame}
  \frametitle{Slide 1}
  Hello Slide 1
\end{frame}
\begin{frame}
  \frametitle{Slide 2}
  Hello Slide 2
\end{frame}
\end{document}
"""
        notes_md = """
# Slide 1: Welcome
Speaker notes for first slide.

# Slide 2: Followup
Speaker notes for second slide.
"""
        with tempfile.TemporaryDirectory() as tmpdir:
            tex_file = os.path.join(tmpdir, "test_notes.tex")
            notes_file = os.path.join(tmpdir, "test_notes.md")
            with open(tex_file, "w") as f:
                f.write(min_tex)
            with open(notes_file, "w") as f:
                f.write(notes_md)

            pdf_file = compile_latex(tex_file, allow_placeholders=True, verbose=False)
            pptx_file = os.path.join(tmpdir, "test_notes.pptx")

            convert_pdf_to_pptx(pdf_file, pptx_file, notes_file=notes_file, verbose=False)

            prs = Presentation(pptx_file)
            self.assertEqual(len(prs.slides), 2)
            self.assertIn("Speaker notes for first slide", prs.slides[0].notes_slide.notes_text_frame.text)
            self.assertIn("Speaker notes for second slide", prs.slides[1].notes_slide.notes_text_frame.text)


if __name__ == "__main__":
    unittest.main()
