"""
Core converter orchestrating PDF extraction and PowerPoint slide reconstruction.
"""

import io
import os
from typing import Optional
import pymupdf
from pptx import Presentation
from pptx.util import Pt

from .icons import add_icons_to_slide, make_icons_only_doc
from .images import extract_image_bytes, make_clipped_images_doc
from .notes import apply_notes_to_slide, load_and_parse_notes
from .shapes import add_drawing_to_slide
from .text import (
    add_paragraphs_to_slide,
    extract_lines_and_icons,
    group_lines_into_paragraphs,
)


def convert_pdf_to_pptx(
    pdf_path: str,
    output_pptx_path: Optional[str] = None,
    notes_file: Optional[str] = None,
    verbose: bool = False,
) -> str:
    """
    Convert a compiled Beamer PDF into an editable PowerPoint (.pptx) presentation,
    faithfully preserving visual layout, colors, typography, shapes, and icons.
    Optionally attaches markdown speaker notes to each slide.
    """
    if not os.path.isfile(pdf_path):
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")

    if not output_pptx_path:
        base = os.path.splitext(pdf_path)[0]
        output_pptx_path = f"{base}.pptx"

    doc = pymupdf.open(pdf_path)
    total_pages = len(doc)
    if total_pages == 0:
        raise ValueError(f"PDF contains 0 pages: {pdf_path}")

    clip_doc = make_clipped_images_doc(pdf_path)
    icon_doc = make_icons_only_doc(pdf_path)

    notes_map = {}
    if notes_file:
        notes_map = load_and_parse_notes(notes_file)
        if verbose:
            print(f"[latex2pptx] Loaded speaker notes for {len(notes_map)} slides from {os.path.basename(notes_file)}")

    prs = Presentation()
    prs.slide_width = Pt(doc[0].rect.width)
    prs.slide_height = Pt(doc[0].rect.height)
    blank_layout = prs.slide_layouts[6]

    applied_notes_count = 0
    for pno in range(total_pages):
        page = doc[pno]
        clip_page = clip_doc[pno]
        icon_page = icon_doc[pno]
        slide = prs.slides.add_slide(blank_layout)

        # Attach speaker notes if present for this slide (1-indexed)
        slide_num = pno + 1
        if slide_num in notes_map:
            apply_notes_to_slide(slide, notes_map[slide_num])
            applied_notes_count += 1

        if verbose and (pno % 5 == 0 or pno == total_pages - 1):
            print(f"[latex2pptx] Processing slide {pno + 1}/{total_pages}...")

        # Z-order drawings and images using bboxlog
        bboxlog = page.get_bboxlog()
        img_seqnos = [i for i, (kind, _) in enumerate(bboxlog) if kind == "fill-image"]
        imgs_info = page.get_image_info(xrefs=True)
        clip_rects = [
            d["scissor"]
            for d in page.get_drawings(extended=True)
            if d.get("type") == "clip" and d.get("scissor") is not None
        ]

        visual_items = []
        for d in page.get_drawings():
            visual_items.append((d.get("seqno", 0), "drawing", d))

        for idx, info in enumerate(imgs_info):
            seq = img_seqnos[idx] if idx < len(img_seqnos) else 1
            visual_items.append((seq, "image", info))

        visual_items.sort(key=lambda x: x[0])

        shapes_rects = []
        for seq, kind, payload in visual_items:
            if kind == "drawing":
                add_drawing_to_slide(slide, payload, page.rect)
                shapes_rects.append(payload["rect"])
            elif kind == "image":
                png_bytes, bbox = extract_image_bytes(doc, clip_page, payload, clip_rects)
                if bbox.width > 0.5 and bbox.height > 0.5:
                    slide.shapes.add_picture(
                        io.BytesIO(png_bytes),
                        Pt(bbox.x0),
                        Pt(bbox.y0),
                        Pt(bbox.width),
                        Pt(bbox.height),
                    )

        # Extract text and icons
        lines, icons = extract_lines_and_icons(page)

        # Add isolated icons with tight alpha crop
        add_icons_to_slide(slide, icon_page, icons, page.rect)

        # Add native editable text boxes
        paragraphs = group_lines_into_paragraphs(lines, shapes_rects)
        add_paragraphs_to_slide(slide, paragraphs, page.rect)

    if notes_file and verbose:
        print(f"[latex2pptx] Successfully attached speaker notes to {applied_notes_count}/{total_pages} slides.")

    prs.save(output_pptx_path)
    return output_pptx_path
