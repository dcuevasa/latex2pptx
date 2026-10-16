"""
Vector shape conversion and geometry classification for PowerPoint presentations.
"""

from typing import Any, Dict, Optional, Tuple
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU_PER_PT = 12700


def rgb_from_float(c: Optional[Tuple[float, ...]]) -> Optional[RGBColor]:
    """Convert a float (r, g, b) tuple (0.0 to 1.0) to an RGBColor."""
    if not c or len(c) < 3:
        return None
    return RGBColor(
        max(0, min(255, int(round(c[0] * 255)))),
        max(0, min(255, int(round(c[1] * 255)))),
        max(0, min(255, int(round(c[2] * 255)))),
    )


def is_axis_aligned_rect(items: list, rect: Any) -> bool:
    """Check if drawing path is an axis-aligned rectangle."""
    if len(items) == 1 and items[0][0] == "re":
        return True
    if len(items) == 4 and all(it[0] == "l" for it in items):
        for _, p1, p2 in items:
            if not (abs(p1.x - p2.x) < 0.15 or abs(p1.y - p2.y) < 0.15):
                return False
        return True
    return False


def is_axis_aligned_rounded_rect(items: list, rect: Any) -> Tuple[bool, float]:
    """Check if drawing path is an axis-aligned rounded rectangle and return corner radius."""
    if len(items) != 8:
        return False, 0.0
    types = [it[0] for it in items]
    if types not in (["l", "c", "l", "c", "l", "c", "l", "c"], ["c", "l", "c", "l", "c", "l", "c", "l"]):
        return False, 0.0
    for it in items:
        if it[0] == "l":
            p1, p2 = it[1], it[2]
            if not (abs(p1.x - p2.x) < 0.15 or abs(p1.y - p2.y) < 0.15):
                return False, 0.0
    for it in items:
        if it[0] == "c":
            p1, p4 = it[1], it[4]
            r = 0.5 * (abs(p4.x - p1.x) + abs(p4.y - p1.y))
            return True, r
    return False, 0.0


def is_ellipse(items: list, rect: Any) -> bool:
    """Check if drawing path is a circle or ellipse (4 cubic beziers)."""
    return len(items) == 4 and all(it[0] == "c" for it in items)


def apply_fill_and_stroke(shape: Any, d: Dict[str, Any]) -> None:
    """Apply fill color, stroke color, and remove default theme drop shadows."""
    fill = d.get("fill")
    fill_op = d.get("fill_opacity", 1.0)
    if fill is not None and (fill_op is None or fill_op > 0.01):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb_from_float(fill)
        if fill_op is not None and fill_op < 0.99:
            spPr = shape._element.spPr
            srgb = spPr.find(f".//{{{A_NS}}}solidFill/{{{A_NS}}}srgbClr")
            if srgb is not None:
                etree.SubElement(srgb, f"{{{A_NS}}}alpha", val=str(int(round(fill_op * 100000))))
    else:
        shape.fill.background()

    stroke = d.get("color")
    width = d.get("width")
    stroke_op = d.get("stroke_opacity", 1.0)
    if stroke is not None and width is not None and width > 0.01 and (stroke_op is None or stroke_op > 0.01):
        shape.line.color.rgb = rgb_from_float(stroke)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()

    # Strip default theme style (removes theme drop shadows and accent borders)
    style = shape._element.find(f"{{{P_NS}}}style")
    if style is not None:
        shape._element.remove(style)

    # Ensure empty effectLst so no shadow is rendered
    spPr = shape._element.spPr
    eff = spPr.find(f"{{{A_NS}}}effectLst")
    if eff is None:
        etree.SubElement(spPr, f"{{{A_NS}}}effectLst")


def add_custgeom_shape(slide: Any, d: Dict[str, Any]) -> Any:
    """Create a native PowerPoint vector shape with custom DrawingML geometry."""
    rect = d["rect"]
    w = max(0.5, rect.width)
    h = max(0.5, rect.height)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Pt(rect.x0),
        Pt(rect.y0),
        Pt(w),
        Pt(h),
    )
    apply_fill_and_stroke(shape, d)

    spPr = shape._element.spPr
    prstGeom = spPr.find(f"{{{A_NS}}}prstGeom")
    if prstGeom is not None:
        spPr.remove(prstGeom)

    custGeom = etree.Element(f"{{{A_NS}}}custGeom")
    etree.SubElement(custGeom, f"{{{A_NS}}}avLst")
    etree.SubElement(custGeom, f"{{{A_NS}}}gdLst")
    etree.SubElement(custGeom, f"{{{A_NS}}}ahLst")
    etree.SubElement(custGeom, f"{{{A_NS}}}cxnLst")
    etree.SubElement(custGeom, f"{{{A_NS}}}rect", l="l", t="t", r="r", b="b")
    pathLst = etree.SubElement(custGeom, f"{{{A_NS}}}pathLst")

    w_emu = max(1, int(round(w * EMU_PER_PT)))
    h_emu = max(1, int(round(h * EMU_PER_PT)))

    has_fill = d.get("fill") is not None
    has_stroke = d.get("color") is not None and d.get("width", 0) > 0.01
    fill_attr = "norm" if has_fill else "none"
    stroke_attr = "1" if has_stroke else "0"

    path_el = etree.SubElement(
        pathLst,
        f"{{{A_NS}}}path",
        w=str(w_emu),
        h=str(h_emu),
        fill=fill_attr,
        stroke=stroke_attr,
    )

    def to_rel(pt: Any) -> Tuple[str, str]:
        return (
            str(int(round((pt.x - rect.x0) * EMU_PER_PT))),
            str(int(round((pt.y - rect.y0) * EMU_PER_PT))),
        )

    cur_pt = None
    for item in d["items"]:
        kind = item[0]
        if kind == "re":
            r = item[1]
            pts = [
                r.tl,
                r.tr,
                r.br,
                r.bl,
            ]
            m_el = etree.SubElement(path_el, f"{{{A_NS}}}moveTo")
            px, py = to_rel(pts[0])
            etree.SubElement(m_el, f"{{{A_NS}}}pt", x=px, y=py)
            for p in pts[1:]:
                l_el = etree.SubElement(path_el, f"{{{A_NS}}}lnTo")
                px, py = to_rel(p)
                etree.SubElement(l_el, f"{{{A_NS}}}pt", x=px, y=py)
            etree.SubElement(path_el, f"{{{A_NS}}}close")
            cur_pt = pts[0]
        elif kind == "l":
            p1, p2 = item[1], item[2]
            if cur_pt is None or abs(cur_pt.x - p1.x) > 0.05 or abs(cur_pt.y - p1.y) > 0.05:
                m_el = etree.SubElement(path_el, f"{{{A_NS}}}moveTo")
                px, py = to_rel(p1)
                etree.SubElement(m_el, f"{{{A_NS}}}pt", x=px, y=py)
            l_el = etree.SubElement(path_el, f"{{{A_NS}}}lnTo")
            px, py = to_rel(p2)
            etree.SubElement(l_el, f"{{{A_NS}}}pt", x=px, y=py)
            cur_pt = p2
        elif kind == "c":
            p1, p2, p3, p4 = item[1], item[2], item[3], item[4]
            if cur_pt is None or abs(cur_pt.x - p1.x) > 0.05 or abs(cur_pt.y - p1.y) > 0.05:
                m_el = etree.SubElement(path_el, f"{{{A_NS}}}moveTo")
                px, py = to_rel(p1)
                etree.SubElement(m_el, f"{{{A_NS}}}pt", x=px, y=py)
            c_el = etree.SubElement(path_el, f"{{{A_NS}}}cubicBezTo")
            for p in (p2, p3, p4):
                px, py = to_rel(p)
                etree.SubElement(c_el, f"{{{A_NS}}}pt", x=px, y=py)
            cur_pt = p4

    if d.get("closePath"):
        etree.SubElement(path_el, f"{{{A_NS}}}close")

    xfrm = spPr.find(f"{{{A_NS}}}xfrm")
    if xfrm is not None:
        spPr.insert(list(spPr).index(xfrm) + 1, custGeom)
    else:
        spPr.insert(0, custGeom)
    return shape


def add_drawing_to_slide(slide: Any, d: Dict[str, Any], page_rect: Any) -> Any:
    """Classify and add vector drawing as the most appropriate native PowerPoint shape."""
    rect = d["rect"]
    # Skip full-page background canvas rectangle at seqno 0
    if (
        d.get("seqno") == 0
        and rect.width >= page_rect.width * 0.98
        and rect.height >= page_rect.height * 0.98
        and d.get("fill") == (1.0, 1.0, 1.0)
        and d.get("color") is None
    ):
        return None

    items = d["items"]
    if is_axis_aligned_rect(items, rect):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Pt(rect.x0),
            Pt(rect.y0),
            Pt(max(0.5, rect.width)),
            Pt(max(0.5, rect.height)),
        )
        apply_fill_and_stroke(shape, d)
        return shape

    is_rrect, radius = is_axis_aligned_rounded_rect(items, rect)
    if is_rrect:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Pt(rect.x0),
            Pt(rect.y0),
            Pt(max(0.5, rect.width)),
            Pt(max(0.5, rect.height)),
        )
        min_dim = max(1.0, min(rect.width, rect.height))
        shape.adjustments[0] = min(0.5, max(0.0, radius / min_dim))
        apply_fill_and_stroke(shape, d)
        return shape

    if is_ellipse(items, rect):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Pt(rect.x0),
            Pt(rect.y0),
            Pt(max(0.5, rect.width)),
            Pt(max(0.5, rect.height)),
        )
        apply_fill_and_stroke(shape, d)
        return shape

    return add_custgeom_shape(slide, d)
